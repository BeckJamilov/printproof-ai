"""Generate the public demo workbook: a FULLY SYNTHETIC capital investment plan.

Every company name, project, manager, amount and comment is invented by this script (fixed random seed,
so the file is reproducible). Nothing is copied from any real workbook.

    python demo/make_synthetic_demo.py demo/Synthetic_Capital_Plan_2027.xlsx

Structure (what Excel Print Optimizer has to handle):
  * a merged title, a subtitle and the reporting unit ("USD thousands") above a column that is usually not printed;
  * two header rows (group headings merged over months, quarters and actuals) and frozen panes;
  * 43 columns: project data, 12 monthly plans, quarterly and annual plan formulas, 12 monthly actuals,
    actual YTD, full-year forecast, variance, completion %, hidden prior-year and helper columns, comments;
  * three levels: segment (1) > programme (1.1) > project (1.1.1), with Excel outline levels;
  * SUBTOTAL(9, ...) formulas on programme and segment rows and the grand total (SUBTOTAL ignores nested
    SUBTOTALs, so every total equals the sum of its projects);
  * inactive projects (zero plan and forecast), some of them with a comment; a programme with no activity;
  * cancelled projects hidden in the workbook; a few Excel notes on cells.
"""
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

random.seed(20270101)
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "demo/Synthetic_Capital_Plan_2027.xlsx")

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
INFO = ["No.", "Project", "Business unit", "Category", "Priority", "Project manager", "Contract status"]   # A..G
PLAN_M = 8                      # H..S   monthly plan
Q1 = PLAN_M + 12                # T..W   quarterly plan (formulas)
ANNUAL = Q1 + 4                 # X      annual plan
ACT_M = ANNUAL + 1              # Y..AJ  monthly actual
YTD = ACT_M + 12                # AK     actual YTD
FCST = YTD + 1                  # AL     full-year forecast
VAR = FCST + 1                  # AM     variance vs plan
COMPL = VAR + 1                 # AN     completion %
PRIOR = COMPL + 1               # AO     prior-year actual (hidden)
HELPER = PRIOR + 1              # AP     helper code (hidden)
COMMENT = HELPER + 1            # AQ     comments
NCOL = COMMENT
NUM_COLS = list(range(PLAN_M, FCST + 1)) + [PRIOR]          # columns that get SUBTOTALs

SEGMENTS = {
    "UPSTREAM — HARBOR BASIN FIELDS": ["Development drilling", "Well workovers", "Field facilities"],
    "MIDSTREAM — PIPELINES AND STORAGE": ["Pipeline integrity", "Storage terminals"],
    "POWER AND UTILITIES": ["Gas-fired power unit", "Grid connections", "Water treatment"],
    "INFRASTRUCTURE AND LOGISTICS": ["Roads and bridges", "Camps and offices"],
    "DIGITAL AND SAFETY": ["IT systems", "Safety and environment"],
    "CORPORATE AND OTHER": ["Studies and engineering", "Legacy claims"],
}
UNITS = ["Harbor Basin", "Cedar Point", "Silver Creek", "Eastgate", "Corporate"]
PRIORITY = ["High", "Medium", "Low"]
STATUS = ["Signed", "Tender", "Planned", "In negotiation"]
ACTIONS = ["Upgrade of", "New", "Replacement of", "Expansion of", "Study for", "Overhaul of"]
OBJECTS = {
    "Development drilling": ["well HB-{n}", "well pad P-{n}", "sidetrack HB-{n}S"],
    "Well workovers": ["workover rig campaign {n}", "artificial lift at HB-{n}", "wellhead set {n}"],
    "Field facilities": ["gathering line GL-{n}", "separator train {n}", "flare system {n}"],
    "Pipeline integrity": ["pipeline section PL-{n}", "cathodic protection zone {n}", "inline inspection run {n}"],
    "Storage terminals": ["tank T-{n}", "loading bay {n}", "metering skid {n}"],
    "Gas-fired power unit": ["turbine unit G{n}", "heat recovery boiler {n}", "control room {n}"],
    "Grid connections": ["substation SS-{n}", "overhead line OL-{n}", "transformer bank {n}"],
    "Water treatment": ["water plant module {n}", "injection pump set {n}", "reservoir {n}"],
    "Roads and bridges": ["access road R-{n}", "bridge B-{n}", "drainage works {n}"],
    "Camps and offices": ["field camp block {n}", "regional office {n}", "warehouse {n}"],
    "IT systems": ["ERP module {n}", "network site {n}", "data centre rack {n}"],
    "Safety and environment": ["gas detection zone {n}", "fire water ring {n}", "emission monitor {n}"],
    "Studies and engineering": ["feasibility study {n}", "front-end design {n}", "survey {n}"],
    "Legacy claims": ["contract claim {n}", "land settlement {n}", "warranty case {n}"],
}
COMMENTS = [
    "Contractor mobilised; schedule confirmed with operations.",
    "Long-lead equipment ordered; delivery expected in Q3.",
    "Tender re-issued after the first round; award planned for Q2.",
    "Scope reduced after technical review; savings reallocated.",
    "Permit pending; spending moved to the second half of the year.",
    "Carry-over from 2026; final invoices only.",
    "Design approved; construction starts after the rainy season.",
    "Cost estimate updated for steel prices (+6%).",
]
INACTIVE_COMMENTS = [
    "On hold until the 2028 budget review.",
    "Deferred: waiting for partner approval.",
]

thin, thick = Side(style="thin"), Side(style="medium")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL_SEG, FILL_PROG, FILL_TOTAL, FILL_HEAD = (PatternFill("solid", fgColor=c) for c in
                                              ("FFB4A7D6", "FFD9D2E9", "FFFFF2CC", "FFE7E6EF"))
NUMFMT = '#,##0.0;-#,##0.0;"-"'

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Capital Plan 2027"

# ---- title block ----
ws["A1"] = "AURORA RIDGE ENERGY (FICTIONAL COMPANY) — CAPITAL INVESTMENT PLAN 2027"
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COMMENT)
ws["A1"].font = Font(bold=True, size=16)
ws["A1"].alignment = Alignment(horizontal="center")
ws["A2"] = "Synthetic demonstration data: all names, projects and amounts are invented."
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COMMENT)
ws["A2"].font = Font(italic=True, size=10)
ws["A2"].alignment = Alignment(horizontal="center")
ws.cell(3, COMPL, "USD thousands").font = Font(bold=True, italic=True)      # unit above column AN
ws.cell(3, COMPL).alignment = Alignment(horizontal="right")

# ---- two header rows ----
groups = [(1, 7, "Project"), (PLAN_M, PLAN_M + 11, "Plan by month"), (Q1, Q1 + 3, "Plan by quarter"),
          (ANNUAL, ANNUAL, "Plan 2027"), (ACT_M, ACT_M + 11, "Actual by month"), (YTD, COMPL, "Performance"),
          (PRIOR, HELPER, "Internal"), (COMMENT, COMMENT, "Comments")]
heads = (INFO + [f"Plan {m}" for m in MONTHS] + ["Q1", "Q2", "Q3", "Q4"] + ["Annual plan"]
         + [f"Actual {m}" for m in MONTHS] + ["Actual YTD", "Full-year forecast", "Variance vs plan",
                                             "Completion, %", "Prior-year actual", "Helper code", "Comments"])
assert len(heads) == NCOL, (len(heads), NCOL)
for a, b, text in groups:
    ws.cell(4, a, text)
    if b > a:
        ws.merge_cells(start_row=4, start_column=a, end_row=4, end_column=b)
for c, h in enumerate(heads, 1):
    ws.cell(5, c, h)
for r in (4, 5):
    for c in range(1, NCOL + 1):
        cell = ws.cell(r, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
        cell.fill = FILL_HEAD
ws.row_dimensions[5].height = 32


def project_values(kind: str):
    """kind: active / inactive / inactive_note / late (no plan, but a forecast) / cancelled."""
    if kind in ("inactive", "inactive_note", "cancelled"):
        return [0.0] * 12, [0.0] * 12, 0.0, round(random.uniform(0, 300), 1) if kind == "cancelled" else 0.0
    if kind == "late":
        plan = [0.0] * 12
        act = [0.0] * 9 + [round(random.uniform(5, 40), 1) for _ in range(3)]
        return plan, act, round(sum(act) + random.uniform(10, 60), 1), 0.0
    size = random.choice([20, 45, 90, 160])
    start = random.randint(0, 5)
    plan = [round(random.uniform(0.4, 1.6) * size, 1) if m >= start else 0.0 for m in range(12)]
    act = [round(p * random.uniform(0.55, 1.15), 1) for p in plan[:9]] + [0.0] * 3      # actuals Jan–Sep
    forecast = round(sum(act) + sum(plan[9:]) * random.uniform(0.8, 1.2), 1)
    return plan, act, forecast, round(random.uniform(0, 1.5) * sum(plan), 1)


def style_row(r: int, fill=None, bold=False, size=None):
    for c in range(1, NCOL + 1):
        cell = ws.cell(r, c)
        cell.border = BOX
        if fill:
            cell.fill = fill
        if bold or size:
            cell.font = Font(bold=bold, size=size or 11)
        if PLAN_M <= c <= PRIOR and c != COMPL:
            cell.number_format = NUMFMT
    ws.cell(r, COMPL).number_format = "0.0%"
    ws.cell(r, COMMENT).alignment = Alignment(wrap_text=True, vertical="top")


def row_formulas(r: int):
    for q in range(4):
        ws.cell(r, Q1 + q, f"=SUM({L(PLAN_M + 3 * q)}{r}:{L(PLAN_M + 3 * q + 2)}{r})")
    ws.cell(r, ANNUAL, f"=SUM({L(Q1)}{r}:{L(Q1 + 3)}{r})")
    ws.cell(r, YTD, f"=SUM({L(ACT_M)}{r}:{L(ACT_M + 11)}{r})")
    ws.cell(r, VAR, f"={L(FCST)}{r}-{L(ANNUAL)}{r}")
    ws.cell(r, COMPL, f"=IF({L(ANNUAL)}{r}=0,0,{L(YTD)}{r}/{L(ANNUAL)}{r})")


def subtotal_row(r: int, first: int, last: int):
    for c in NUM_COLS:
        ws.cell(r, c, f"=SUBTOTAL(9,{L(c)}{first}:{L(c)}{last})")
    ws.cell(r, VAR, f"={L(FCST)}{r}-{L(ANNUAL)}{r}")
    ws.cell(r, COMPL, f"=IF({L(ANNUAL)}{r}=0,0,{L(YTD)}{r}/{L(ANNUAL)}{r})")


row = 6
first_body = row
serial = 100
notes = 0
for s, (segment, programmes) in enumerate(SEGMENTS.items(), 1):
    seg_row = row
    ws.cell(row, 1, str(s))
    ws.cell(row, 2, segment)
    row += 1
    for p, prog in enumerate(programmes, 1):
        prog_row = row
        ws.cell(row, 1, f"{s}.{p}")
        ws.cell(row, 2, prog)
        ws.row_dimensions[row].outline_level = 1
        row += 1
        dormant = prog in ("Legacy claims", "Camps and offices")          # whole programme without activity
        for i in range(1, random.randint(7, 13) + 1):
            serial += random.randint(1, 9)
            roll = random.random()
            kind = ("inactive" if dormant else "cancelled" if roll < 0.10 else "inactive" if roll < 0.40
                    else "inactive_note" if roll < 0.45 else "late" if roll < 0.50 else "active")
            plan, act, forecast, prior = project_values(kind)
            ws.cell(row, 1, f"{s}.{p}.{i}")
            obj = random.choice(OBJECTS[prog]).format(n=serial)
            action = random.choice(ACTIONS)
            ws.cell(row, 2, f"{action} {obj}" if action != "New" else f"New {obj}")
            ws.cell(row, 3, random.choice(UNITS))
            ws.cell(row, 4, "Opex" if prog in ("Well workovers", "Studies and engineering") else "Capex")
            ws.cell(row, 5, random.choice(PRIORITY))
            ws.cell(row, 6, f"Manager {random.randint(1, 24):02d}")
            ws.cell(row, 7, random.choice(STATUS))
            for m in range(12):
                ws.cell(row, PLAN_M + m, plan[m])
                ws.cell(row, ACT_M + m, act[m])
            ws.cell(row, FCST, forecast)
            ws.cell(row, PRIOR, prior)
            ws.cell(row, HELPER, f"WBS-{s}{p:02d}-{serial}")
            row_formulas(row)
            if kind == "active" and random.random() < 0.45:
                ws.cell(row, COMMENT, random.choice(COMMENTS))
            if kind == "inactive_note":
                ws.cell(row, COMMENT, random.choice(INACTIVE_COMMENTS))
            if kind == "cancelled":
                ws.cell(row, COMMENT, "Cancelled in 2026; kept for the audit trail.")
                ws.row_dimensions[row].hidden = True
            if kind == "active" and notes < 4 and random.random() < 0.15:
                ws.cell(row, FCST).comment = Comment("Forecast includes the approved change order.", "PrintProof AI demo")
                notes += 1
            ws.row_dimensions[row].outline_level = 2
            style_row(row)
            row += 1
        subtotal_row(prog_row, prog_row + 1, row - 1)
        style_row(prog_row, FILL_PROG, bold=True)
    subtotal_row(seg_row, seg_row + 1, row - 1)
    style_row(seg_row, FILL_SEG, bold=True, size=12)
grand = row
ws.cell(grand, 2, "GRAND TOTAL")
subtotal_row(grand, first_body, grand - 1)
style_row(grand, FILL_TOTAL, bold=True, size=12)

widths = {1: 7, 2: 40, 3: 13, 4: 9, 5: 9, 6: 12, 7: 13, ANNUAL: 11, YTD: 11, FCST: 12, VAR: 11, COMPL: 11,
          PRIOR: 11, HELPER: 15, COMMENT: 42}
for c in range(1, NCOL + 1):
    ws.column_dimensions[L(c)].width = widths.get(c, 9)
for c in (PRIOR, HELPER):
    ws.column_dimensions[L(c)].hidden = True
ws.freeze_panes = "H6"
ws.print_area = f"A1:{L(COMMENT)}{grand}"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "portrait"
wb.properties.creator = "PrintProof AI demo"
wb.properties.lastModifiedBy = "PrintProof AI demo"
wb.properties.title = "Synthetic capital investment plan (fictional data)"
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"{OUT}: {grand} rows, {NCOL} columns, print area {ws.print_area}")
