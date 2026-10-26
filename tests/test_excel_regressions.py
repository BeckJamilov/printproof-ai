"""Excel Print Optimizer — regression tests for the bugs found by the automated QA.

Pure tests (no Excel): document metadata classification, frozen header rows, balanced page breaks,
nested text numbering, clipped-text detection, Readability First page priority.
Excel tests (skipped without Microsoft Excel): small synthetic workbooks generated at test time.
"""

import hashlib
import sys
from pathlib import Path

import pymupdf
import pytest

PROJECT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_DIR))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from test_excel_print import SETTINGS, FakeEngine, cell, make_sheet  # noqa: E402

from excel_print.columns import index_to_letter  # noqa: E402
from excel_print.excel_engine import excel_available  # noqa: E402
from excel_print.header_meta import MetaCell, plan_metadata  # noqa: E402
from excel_print.layout import balanced_starts  # noqa: E402
from excel_print.model import OptimizeSettings  # noqa: E402
from excel_print.numbering import renumber  # noqa: E402
from excel_print.optimizer import _numbering_column, optimize  # noqa: E402
from excel_print.row_filter import column_roles  # noqa: E402
from excel_print.verify import hidden_text  # noqa: E402
from excel_print.workbook_info import frozen_rows  # noqa: E402


# ---- 1. document metadata in unprinted columns -------------------------------------------------------------

HEADER = [3, 4, 5, 6]          # 3 title, 4 units, 5-6 column headings
PRINTED = [4, 5, 10, 75, 78]   # D, E, J, BW, BZ
TEXT = {(3, 4): "FORECAST OF THE PROGRAMME", (5, 4): "No.", (5, 5): "Name", (5, 78): "Comments", (6, 10): "Plan"}
MERGES = {(3, c): "$D$3:$BZ$3" for c in PRINTED}


def test_unit_label_in_an_unprinted_column_is_carried_to_the_nearest_free_printed_cell():
    plan = plan_metadata(HEADER, PRINTED, TEXT, MERGES, [MetaCell(4, 77, "тыс. $")])        # BY4
    assert plan.metadata_rows == [3, 4]
    assert [(p.source, p.target, p.text) for p in plan.placements] == [((4, 77), (4, 78), "тыс. $")]


def test_column_headings_hidden_helpers_and_bare_numbers_are_not_carried_over():
    cands = [MetaCell(5, 77, "Execution"),                 # a column heading (row 5 is a heading row)
             MetaCell(4, 70, "helper code", hidden_col=True),
             MetaCell(4, 71, "12345.67"),
             MetaCell(4, 72, "x" * 120)]
    plan = plan_metadata(HEADER, PRINTED, TEXT, MERGES, cands)
    assert plan.placements == []
    assert len(plan.skipped) == 3 and any("hidden" in s for s in plan.skipped)


def test_reporting_dates_are_metadata_but_numbers_are_not():
    plan = plan_metadata(HEADER, PRINTED, TEXT, MERGES, [MetaCell(4, 60, "31.07.2022")])
    assert plan.placements and plan.placements[0].text == "31.07.2022"


def test_title_anchored_in_an_unprinted_column_is_visible_through_its_merged_area():
    text = {k: v for k, v in TEXT.items() if k != (3, 4)}
    plan = plan_metadata(HEADER, [5, 10, 75, 78], text, {(3, c): "$D$3:$BZ$3" for c in (5, 10, 75, 78)},
                         [MetaCell(3, 4, "FORECAST OF THE PROGRAMME", merged=True, merge_cols=(4, 78))])
    assert plan.placements == [] and [m.text for m in plan.visible_merged] == ["FORECAST OF THE PROGRAMME"]


def test_no_free_printed_cell_is_reported_not_forced():
    merged_row4 = {**MERGES, **{(4, c): "$D$4:$BZ$4" for c in PRINTED}}     # row 4 is one merged (empty) band
    plan = plan_metadata(HEADER, PRINTED, TEXT, merged_row4, [MetaCell(4, 77, "USD thousand")])
    assert plan.placements == [] and "no free printed cell" in plan.skipped[0]


# ---- 2. header rows from the frozen panes in the file -------------------------------------------------------

def test_frozen_rows_are_read_from_the_workbook_file(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "x"
    wb.active.freeze_panes = "A4"
    wb.save(tmp_path / "f.xlsx")
    assert frozen_rows(tmp_path / "f.xlsx", "S") == (1, 3)
    wb.active.freeze_panes = None
    wb.save(tmp_path / "g.xlsx")
    assert frozen_rows(tmp_path / "g.xlsx", "S") is None


# ---- 3. balanced page breaks ------------------------------------------------------------------------------

def test_balanced_breaks_avoid_an_isolated_last_row():
    heights = [10.0] * 31                              # greedy: 30 rows + 1 row alone on page 2
    starts = balanced_starts(heights, [False] * 31, 300, 300, 2)
    assert starts and 14 <= starts[0] <= 17


def test_balanced_breaks_keep_a_heading_with_the_next_row():
    heights = [10.0] * 20
    keep = [False] * 20
    keep[9] = True                                      # row 9 is a section heading
    starts = balanced_starts(heights, keep, 150, 150, 2)
    assert starts and starts[0] != 10                   # never break between the heading and its first row


def test_balanced_breaks_report_impossible_splits():
    assert balanced_starts([10.0] * 50, [False] * 50, 100, 100, 2) is None


# ---- 4. nested numbering stored as text -----------------------------------------------------------------------

def test_text_nested_numbers_are_recognised_and_renumbered_hierarchically():
    rows = [{"c": cell("1", text="1"), "label": cell("TRUCKS")}, {"c": cell("1.1", text="1.1")}, {"c": cell("1.2", text="1.2")},
            {"c": cell("1.4", text="1.4")}, {"c": cell("3", text="3"), "label": cell("CRANES")}, {"c": cell("3.1", text="3.1")},
            {"c": cell("3.3", text="3.3")}]
    sheet = make_sheet([{"label": cell("Header")}] + rows)
    body = [r for r in sheet.rows if r.number > 1]
    roles = column_roles(sheet, body)
    assert roles[3] == "index"
    assert _numbering_column(sheet, roles, body) == 3
    new = renumber([(r.number, r.cell(3).text) for r in body])
    printed = [new.get(r.number, (r.cell(3).text,))[0] for r in body]
    assert printed == ["1", "1.1", "1.2", "1.3", "2", "2.1", "2.2"]


def test_years_or_codes_are_not_treated_as_a_sequence_numbering():
    rows = [{"c": cell(float(y), text=str(y))} for y in (2019, 2020, 2021, 2022, 2023)]
    sheet = make_sheet([{"label": cell("Header")}] + rows)
    body = [r for r in sheet.rows if r.number > 1]
    assert _numbering_column(sheet, column_roles(sheet, body), body) is None


# ---- 5. clipped text detection in the rendered PDF -----------------------------------------------------------------

def _page_with_cover(cover):
    doc = pymupdf.open()
    page = doc.new_page(width=842, height=595)
    page.insert_text((80, 100), "FORECAST OF THE PROGRAMME", fontsize=16)
    page.insert_text((80, 200), "lowercase only words", fontsize=12)
    if cover:
        page.draw_rect(pymupdf.Rect(70, 80, 600, 92), color=None, fill=(1, 1, 1))   # hides the top of the capitals
    return pymupdf.open(stream=doc.tobytes(), filetype="pdf")[0]


def test_text_cut_at_the_top_of_its_cell_is_detected():
    assert hidden_text(_page_with_cover(True)) == ["FORECAST OF THE PROGRAMME"]


def test_intact_text_including_lowercase_only_is_not_flagged():
    assert hidden_text(_page_with_cover(False)) == []


# ---- 6. Readability First: never more pages than at the minimum font ------------------------------------------------

def test_readability_never_trades_a_page_for_larger_text(tmp_path):
    # 1 page up to 90 %, 2 pages above: the minimum (10 pt = 84 %) fits on 1 page, so the result must be 1 page
    res = optimize(tmp_path / "x.xlsx", "S", OptimizeSettings(columns_text="B, D, E", filter_columns_text="D, E",
                                                             mode="readability", min_font_pt=10),
                   engine_factory=lambda p: FakeEngine(p, fit_zoom=90, row_h=400.0))
    assert res.success and res.summary["pages"] == 1 and res.summary["zoom_percent"] == 90


def test_extremely_small_text_is_reported_as_a_readability_warning(tmp_path):
    res = optimize(tmp_path / "x.xlsx", "S", OptimizeSettings(columns_text="B, D, E", filter_columns_text="D, E",
                                                             mode="pages", max_pages=1),
                   engine_factory=lambda p: FakeEngine(p, fit_zoom=30, row_h=400.0))      # 12 pt x 30 % = 3.6 pt
    assert res.summary["readability"] == "extremely small"
    assert "extremely small" in res.message and any("extremely small" in w for w in res.warnings)


# ---- 7. Excel end-to-end regressions (synthetic workbooks) ------------------------------------------------------------

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

needs_excel = pytest.mark.skipif(not excel_available()[0], reason="Microsoft Excel is not installed")
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _report(path: Path, items: int = 40, unit_col: str = "D", hidden_helper: bool = True) -> None:
    """A budget report: merged title, unit label in `unit_col` row 2, headings in row 3 (frozen),
    sections with nested text numbers, zero items, a hidden helper column G, SUBTOTAL totals."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "QUARTERLY CAPITAL EXPENDITURE FORECAST FOR ALL REGIONS — DEMO DATA"
    ws.merge_cells("A1:F1")
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws[f"{unit_col}2"] = "EUR thousand"
    for c, h in enumerate(["No.", "Item", "Plan", "Actual", "Variance", "Note", "Helper"], 1):
        ws.cell(3, c, h).font = Font(bold=True)
        ws.cell(3, c).border = BOX
    row, numbers = 4, []
    for sec in (1, 2, 3):
        ws.cell(row, 1, str(sec))
        ws.cell(row, 2, f"SECTION {sec}")
        sec_row, first = row, row + 1
        row += 1
        for i in range(1, items // 3 + 1):
            zero = sec == 3 or i % 4 == 0
            ws.cell(row, 1, f"{sec}.{i}")
            ws.cell(row, 2, f"Item {sec}-{i} with a reasonably long description of the work")
            ws.cell(row, 3, 0 if zero else 100 + i)
            ws.cell(row, 4, 0 if zero else 90 + i)
            ws.cell(row, 5, f"=D{row}-C{row}")
            ws.cell(row, 7, f"HELPER-{sec}-{i}")
            for c in range(1, 8):
                ws.cell(row, c).border = BOX
            row += 1
        for col in "CDE":
            ws[f"{col}{sec_row}"] = f"=SUBTOTAL(9,{col}{first}:{col}{row - 1})"
        for c in range(1, 8):
            ws.cell(sec_row, c).border = BOX
            ws.cell(sec_row, c).fill = PatternFill("solid", fgColor="FFDDEBF7")
            ws.cell(sec_row, c).font = Font(bold=True)
    ws.cell(row, 2, "GRAND TOTAL")
    for col in "CDE":
        ws[f"{col}{row}"] = f"=SUBTOTAL(9,{col}5:{col}{row - 1})"
    for c in range(1, 8):
        ws.cell(row, c).border = BOX
    for col, w in zip("ABCDEFG", (6, 30, 10, 10, 10, 20, 12)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["G"].hidden = hidden_helper
    ws.print_area = f"A1:G{row}"
    ws.freeze_panes = "A4"
    wb.save(path)


def _pdf_text(pdf: bytes) -> list[str]:
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        return [" ".join(p.get_text().split()) for p in doc]


@pytest.fixture(scope="module")
def report_runs(tmp_path_factory):
    d = tmp_path_factory.mktemp("reg")
    path = d / "report.xlsx"
    _report(path, items=90)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    one = optimize(path, "Report", OptimizeSettings(columns_text="A, B, C, E", filter_columns_text="C, D",
                                                    mode="pages", max_pages=1, paper="A4", orientation="portrait"))
    big = optimize(path, "Report", OptimizeSettings(columns_text="A, B, E", filter_columns_text="C, D",
                                                    mode="readability", min_font_pt=12, paper="A4", orientation="portrait"))
    return {"one": one, "big": big, "unchanged": hashlib.sha256(path.read_bytes()).hexdigest() == before}


@needs_excel
def test_unit_label_survives_when_its_column_is_not_printed(report_runs):
    res = report_runs["one"]
    assert res.success, res.message
    assert "EUR thousand" in " ".join(_pdf_text(res.pdf_bytes))            # D2 is not printed (D only filters)
    assert res.summary["metadata_carried_over"]


@needs_excel
def test_hidden_helper_column_never_leaks(report_runs):
    for res in (report_runs["one"], report_runs["big"]):
        assert "HELPER-" not in " ".join(_pdf_text(res.pdf_bytes))


@needs_excel
def test_merged_title_is_complete_and_visible_on_a_narrow_table(report_runs):
    res = report_runs["big"]                    # only A, B, E printed at 12 pt on A4 portrait: the title must wrap
    assert res.success, res.message
    text = _pdf_text(res.pdf_bytes)[0].replace(" ", "")
    assert "QUARTERLYCAPITALEXPENDITUREFORECASTFORALLREGIONS—DEMODATA" in text
    with pymupdf.open(stream=res.pdf_bytes, filetype="pdf") as doc:
        assert hidden_text(doc[0]) == []


@needs_excel
def test_column_headings_repeat_on_every_page(report_runs):
    res = report_runs["big"]
    pages = _pdf_text(res.pdf_bytes)
    assert len(pages) >= 2
    assert all("Item" in p and "Variance" in p for p in pages)


@needs_excel
def test_nested_text_numbers_are_consecutive_after_filtering(report_runs):
    res = report_runs["one"]
    printed = [d.printed_number for d in res.decisions if d.keep and d.printed_number]
    assert printed                                   # gaps (x.4, x.8, …) and the zero section 3 were closed
    kept_numbers = [d.printed_number or "" for d in res.decisions if d.keep and d.category != "header"]
    assert "3" not in [n for n in kept_numbers if n]  # section 3 (all zero) was removed, not renumbered into view


@needs_excel
def test_readability_pages_are_balanced_and_values_preserved(report_runs):
    res = report_runs["big"]
    fills = [c for c in res.checks if c.name.startswith("All selected values")]
    assert fills and fills[0].passed                 # every printed value equals Excel's calculated display text
    rows_per_page = []
    with pymupdf.open(stream=res.pdf_bytes, filetype="pdf") as doc:
        for p in doc:
            rows_per_page.append(sum(1 for w in p.get_text("words") if w[4].startswith("Item")))
    assert min(rows_per_page[1:] or [99]) > 2, rows_per_page    # no page with only one or two rows


@needs_excel
def test_source_workbook_unchanged(report_runs):
    assert report_runs["unchanged"]


def test_text_column_without_body_text_keeps_its_heading_readable():
    from excel_print.optimizer import _natural_widths
    sheet = make_sheet([{"label": cell("Header"), "g": cell("Comment", bold=True)}, {}, {}])   # G: heading only
    natural, minimum = _natural_widths(sheet, {1, 2, 3}, [7])
    assert minimum[7] >= len("Comment") and natural[7] >= minimum[7]
