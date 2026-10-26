"""Excel Print Optimizer — end-to-end tests with the real rendering engine (Microsoft Excel).

Synthetic workbooks are created in a temporary folder at test time (no workbook is stored in Git).
These tests are skipped, not faked, when Microsoft Excel is not available on the machine.
"""

import hashlib
import re
import sys
from pathlib import Path

import pymupdf
import pytest

PROJECT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_DIR))

from excel_print.excel_engine import ExcelEngineError, ExcelSession, excel_available  # noqa: E402
from excel_print.model import OptimizeSettings  # noqa: E402
from excel_print.optimizer import optimize  # noqa: E402

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Border, Font, PatternFill, Side  # noqa: E402

pytestmark = pytest.mark.skipif(not excel_available(), reason="Microsoft Excel (the rendering engine) is not installed")

LONG_COMMENT = ("Contractor confirmed a three-month delay because of permit issues; the remaining budget is moved "
                "to next year and the steering committee will review the revised schedule in March.")


def build_workbook(path: Path, n_projects: int = 30, comment_every: int = 7) -> dict:
    """A small budget sheet: header, two sections, projects, a SUBTOTAL row and a grand total.

    Columns: A no. | B description | C ignored | D plan | E fact | F ignored | G comment.
    Returns the facts the tests check against.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["No.", "Project description", "Internal code", "Plan 2026", "Fact 2026", "Owner", "Comment"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    facts = {"kept_values": [], "filtered_names": [], "comments": [], "sections": []}
    row = 2
    for section in ("Section A — Infrastructure", "Section B — Equipment"):
        ws.cell(row, 2, section).font = Font(bold=True, size=13)
        for c in range(1, 8):
            ws.cell(row, c).fill = PatternFill("solid", fgColor="FFD9E1F2")
        facts["sections"].append(section)
        first = row + 1
        row += 1
        for i in range(n_projects // 2):
            n = row
            ws.cell(row, 1, n)
            ws.cell(row, 3, f"X-{n}")
            ws.cell(row, 6, "Finance")
            if i % 3 == 2:  # every third project has nothing in D and E -> filtered
                name = f"Dormant project {n}"
                ws.cell(row, 4, 0)
            else:
                name = f"Active project {n}"
                ws.cell(row, 4, 1000 * n + 0.5)
                ws.cell(row, 5, 250 * n)
                ws.cell(row, 4).number_format = "#,##0.00"
                ws.cell(row, 5).number_format = "#,##0.00"
                facts["kept_values"].append(f"{1000 * n + 0.5:,.2f}")
            if i % comment_every == 1:
                if name.startswith("Dormant"):
                    name = f"Commented project {n}"
                ws.cell(row, 7, LONG_COMMENT)
                facts["comments"].append(n)
            elif name.startswith("Dormant"):
                facts["filtered_names"].append(name)
            ws.cell(row, 2, name)
            row += 1
        ws.cell(row, 2, "Subtotal")
        ws.cell(row, 4, f"=SUBTOTAL(9,D{first}:D{row - 1})")
        ws.cell(row, 5, f"=SUBTOTAL(9,E{first}:E{row - 1})")
        row += 1
    ws.cell(row, 2, "Grand total").font = Font(bold=True)
    ws.cell(row, 4, f"=SUBTOTAL(9,D2:D{row - 1})")
    ws.cell(row, 2).comment = Comment("Checked by controller", "test")
    facts["last_row"] = row
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["G"].width = 20
    ws.print_area = f"A1:G{row}"
    ws.freeze_panes = "A2"
    wb.save(path)
    return facts


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def number_pattern(formatted: str) -> re.Pattern:
    """'3,000.50' -> matches 3,000.50 / 3 000,50 / 3.000,50 (thousand and decimal separators depend on locale)."""
    whole, decimals = formatted.split(".")
    groups = whole.split(",")
    return re.compile(r"(?<![\d])" + r"[\s  .,']?".join(groups) + r"[.,]" + decimals + r"(?![\d])")


def pdf_text(pdf: bytes) -> str:
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        return " ".join(" ".join(p.get_text().split()) for p in doc)


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    path = tmp_path_factory.mktemp("xl") / "synthetic_budget.xlsx"
    return path, build_workbook(path)


@pytest.fixture(scope="module")
def readable_run(workbook):
    path, facts = workbook
    before = sha(path)
    settings = OptimizeSettings(columns_text="B, D, E, G", filter_columns_text="D, E", min_font_pt=8)
    res = optimize(path, "Budget", settings)
    return res, facts, before, sha(path)


def test_selected_columns_and_values_are_printed_verified(readable_run):
    res, facts, _, _ = readable_run
    assert res.success, (res.message, [(c.name, c.detail) for c in res.checks if c.passed is False])
    text = pdf_text(res.pdf_bytes)
    assert "Internal code" not in text and "X-5" not in text and "Finance" not in text  # columns C and F hidden
    for value in facts["kept_values"]:
        assert number_pattern(value).search(text), value  # preserved values with 2 decimals (any locale)


def test_row_filter_keeps_headings_totals_and_comments(readable_run):
    res, facts, _, _ = readable_run
    text = pdf_text(res.pdf_bytes)
    for name in facts["filtered_names"]:
        assert name not in text
    for section in facts["sections"]:
        assert section in text
    assert text.count("Subtotal") == 2 and "Grand total" in text
    reasons = {d.row: d.category for d in res.decisions}
    assert all(reasons[r] in ("comment", "values") for r in facts["comments"])


def test_long_wrapped_comments_are_complete(readable_run):
    res, facts, _, _ = readable_run
    text = pdf_text(res.pdf_bytes)
    assert text.count(LONG_COMMENT) == len(facts["comments"])
    assert any(c.passed and "nothing clipped" in c.name for c in res.checks)


def test_minimum_font_is_respected_and_measured(readable_run):
    res, _, _, _ = readable_run
    assert res.summary["measured_min_font_pt"] >= 8 - 0.05
    with pymupdf.open(stream=res.pdf_bytes, filetype="pdf") as doc:
        w, h = doc[0].rect.width, doc[0].rect.height
    assert abs(w - 1190.55) < 2 and abs(h - 841.89) < 2  # A3 landscape


def test_original_workbook_is_unchanged(readable_run):
    _, _, before, after = readable_run
    assert before == after


def test_page_limit_mode(workbook):
    path, _ = workbook
    settings = OptimizeSettings(columns_text="B, D, E, G", filter_columns_text="D, E", mode="pages", max_pages=1)
    res = optimize(path, "Budget", settings)
    assert res.success and res.summary["pages"] == 1


def test_impossible_font_and_page_combination_is_explained_not_faked(tmp_path):
    path = tmp_path / "big.xlsx"
    build_workbook(path, n_projects=400, comment_every=2)
    settings = OptimizeSettings(columns_text="B, D, E, G", filter_columns_text="D, E", mode="pages", max_pages=1)
    res = optimize(path, "Budget", settings)
    if res.success:  # Excel fitted it, then the text must be small and the user warned
        assert res.summary["pages"] == 1 and res.warnings
    elif res.pdf_bytes is None:  # the model already shows it is impossible
        assert "cannot fit" in " ".join(res.summary.get("explanation", [res.message]))
    else:  # only a failing draft exists: it is clearly not called a success, and the failures are listed
        assert "did not pass every check" in res.message
        assert any(c.passed is False for c in res.checks)


def test_failed_pdf_conversion_is_reported(workbook, monkeypatch):
    path, _ = workbook

    def broken_export(self, out):
        raise ExcelEngineError("Excel could not export the PDF: simulated printer failure")

    monkeypatch.setattr(ExcelSession, "export_pdf", broken_export)
    res = optimize(path, "Budget", OptimizeSettings(columns_text="B, D, E, G", filter_columns_text="D, E"))
    assert not res.success and res.pdf_bytes is None and "simulated printer failure" in res.message


def test_missing_sheet_and_columns_outside_print_area_are_clear_errors(workbook):
    path, _ = workbook
    res = optimize(path, "Nope", OptimizeSettings(columns_text="B, D", filter_columns_text="D"))
    assert not res.success and "Nope" in res.message
    res = optimize(path, "Budget", OptimizeSettings(columns_text="B, D, K", filter_columns_text="D"))
    assert not res.success and "K" in res.message


# ---- a grouped (outlined) budget report: hierarchy, empty headings, numbering  ----------------------------

GROUPED = {
    # (number, description, fill?, level, plan, forecast, comment)
    "rows": [
        ("1", "PORTFOLIO TOTAL", True, 1, 500.0, 120.0, None),
        ("2", "DRILLING", True, 1, 20.0, 2.5, None),
        ("3", "Drilling north (inactive)", False, 2, 0, 0, None),
        ("4", "Drilling south", False, 2, 20.0, 2.5, LONG_COMMENT),
        ("5", "EMPTY SECTION", True, 1, 0, 0, None),
        ("6", "Nothing here", False, 2, 0, 0, None),
        ("7", "DEVELOPMENT", True, 1, 480.0, 117.5, None),
        ("8", "Objects (coloured heading, empty)", True, 2, 0, 0, None),
        ("9", "Construction programme", False, 2, 0, 0, None),
        ("9.1.", "Main contract (empty)", True, 2, 0, 0, None),
        ("9.2.", "Additional objects", True, 2, 480.0, 115.0, LONG_COMMENT),
        ("12", "Insurance (coloured, empty)", True, 2, 0, 0, None),
        ("13", "Licences", False, 2, 0, 0, "Zero this year, renewal is paid by head office."),
        ("14", "CLAIMS", True, 1, 0, 0, None),
        ("15", "Claim north", False, 2, 0, 0, None),
    ],
}


def build_grouped(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["No.", "Description", "Plan", "Forecast", "Comment"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, (num, desc, fill, level, plan, fc, comment) in enumerate(GROUPED["rows"], start=2):
        ws.cell(i, 1, int(num) if num.isdigit() else num)
        ws.cell(i, 2, desc)
        ws.cell(i, 3, plan).number_format = "#,##0.00"
        ws.cell(i, 4, fc).number_format = "#,##0.00"
        if comment:
            ws.cell(i, 5, comment)
        if fill:
            for c in range(1, 6):
                ws.cell(i, c).fill = PatternFill("solid", fgColor="FFCCC0DA")
                ws.cell(i, c).font = Font(bold=True, size=13 if level == 1 else 11)
        ws.row_dimensions[i].outline_level = level - 1
    thin = Side(style="thin")
    for row in ws.iter_rows(min_row=1, max_row=len(GROUPED["rows"]) + 1, max_col=5):
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["E"].width = 25
    ws.print_area = f"A1:E{len(GROUPED['rows']) + 1}"
    ws.freeze_panes = "A2"
    wb.save(path)


@pytest.fixture(scope="module")
def grouped_run(tmp_path_factory):
    path = tmp_path_factory.mktemp("xlg") / "grouped.xlsx"
    build_grouped(path)
    before = sha(path)
    original_numbers = [c.value for c in openpyxl.load_workbook(path).active["A"]]
    res = optimize(path, "Report", OptimizeSettings(columns_text="A-E", filter_columns_text="C, D",
                                                     mode="pages", max_pages=1))
    after_numbers = [c.value for c in openpyxl.load_workbook(path).active["A"]]
    return res, before == sha(path), original_numbers == after_numbers


def test_grouped_empty_headings_and_sections_are_not_printed(grouped_run):
    res, _, _ = grouped_run
    assert res.success, (res.message, [(c.name, c.detail) for c in res.checks if c.passed is False])
    text = pdf_text(res.pdf_bytes)
    for gone in ("Drilling north", "EMPTY SECTION", "Nothing here", "Objects (coloured", "Main contract",
                 "Insurance", "CLAIMS", "Claim north"):
        assert gone not in text, gone
    for kept in ("PORTFOLIO TOTAL", "DRILLING", "Drilling south", "DEVELOPMENT", "Additional objects", "Licences"):
        assert kept in text, kept


def test_grouped_parent_of_a_relevant_sub_item_is_kept(grouped_run):
    res, _, _ = grouped_run
    by_label = {d.label: d for d in res.decisions}
    assert by_label["Construction programme"].keep and by_label["Construction programme"].category == "parent"
    assert by_label["Additional objects"].keep            # 9.2. is printed under its parent 9, not orphaned
    assert by_label["Licences"].category == "comment"     # zero figures, but a comment


def test_grouped_numbering_is_consecutive_in_the_pdf(grouped_run):
    res, _, _ = grouped_run
    printed = [d.printed_number or str(d.label) for d in res.decisions if d.keep and d.row > 1]
    numbers = [d.printed_number for d in res.decisions if d.keep and d.row > 1 and d.printed_number]
    assert numbers  # the gaps were closed
    with pymupdf.open(stream=res.pdf_bytes, filetype="pdf") as doc:
        words = doc[0].get_text("words")
    desc_left = min(w[0] for w in words if w[4] in ("PORTFOLIO", "DRILLING", "Licences"))
    first_col = [w[4] for w in sorted(words, key=lambda w: w[1]) if w[2] <= desc_left and w[4][0].isdigit()]
    assert first_col == ["1", "2", "3", "4", "5", "5.1.", "6"], (first_col, printed)


def test_grouped_source_workbook_is_unchanged(grouped_run):
    _, same_hash, same_numbers = grouped_run
    assert same_hash and same_numbers   # renumbering happened only in the temporary copy


def test_grouped_layout_wraps_comments_and_gives_text_more_width(grouped_run):
    res, _, _ = grouped_run
    names = {c.name: c for c in res.checks}
    assert names["No text cut off at cell borders"].passed is True
    assert next(c for n, c in names.items() if n.startswith("Safe margins")).passed is True
    w = res.column_widths
    assert w["E"] > w["C"] and w["B"] > w["D"]   # comment and description wider than the number columns
    assert pdf_text(res.pdf_bytes).count(LONG_COMMENT) == 2


def test_grouped_one_page_scale_is_the_largest_that_fits(grouped_run):
    res, _, _ = grouped_run
    z = res.summary["zoom_percent"]
    tried = " ".join(res.attempts)
    assert res.summary["pages"] == 1
    # either the next larger scale was tried in a real PDF and failed, or the page width is the limit
    assert f"{z + 1} %: 2 page(s)" in tried or "page width" in " ".join(res.summary["explanation"]) or z == 120
